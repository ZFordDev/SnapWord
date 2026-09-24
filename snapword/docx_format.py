"""Basic rich DOCX interchange. Advanced Word layout is explicitly unsupported."""

import base64
import io
from html import escape

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.table import Table
from docx.text.run import Run
from PySide6.QtCore import Qt
from PySide6.QtGui import QTextCharFormat, QTextDocument, QTextListFormat, QTextTable
from PySide6.QtWidgets import QApplication


def _run_html(run, paragraph):
    styles = []
    fonts = (run.font, run.style.font, paragraph.style.font)
    for attr, css, value in (
        ("bold", "font-weight", "700"),
        ("italic", "font-style", "italic"),
        ("underline", "text-decoration", "underline"),
        ("strike", "text-decoration", "line-through"),
    ):
        setting = next((getattr(font, attr) for font in fonts if getattr(font, attr) is not None), False)
        if setting:
            styles.append(f"{css}:{value}")
    if run.font.name:
        styles.append(f"font-family:{escape(run.font.name, quote=True)}")
    if run.font.size:
        styles.append(f"font-size:{run.font.size.pt}pt")
    if run.font.color.rgb:
        styles.append(f"color:#{run.font.color.rgb}")
    if run.font.superscript:
        styles.append("vertical-align:super")
    if run.font.subscript:
        styles.append("vertical-align:sub")
    content = escape(run.text).replace("\n", "<br>")
    for blip in run._r.xpath(".//a:blip"):
        relationship = blip.get(qn("r:embed"))
        if relationship and relationship in run.part.rels:
            part = run.part.rels[relationship].target_part
            data = base64.b64encode(part.blob).decode("ascii")
            content += f'<img src="data:{part.content_type};base64,{data}">'
    return f'<span style="{";".join(styles)}">{content}</span>'


def _paragraph_html(paragraph):
    parts = []
    for item in paragraph.iter_inner_content():
        if isinstance(item, Run):
            parts.append(_run_html(item, paragraph))
        else:
            text = "".join(_run_html(run, paragraph) for run in item.runs)
            parts.append(f'<a href="{escape(item.url, quote=True)}">{text}</a>')
    tag = "p"
    style = paragraph.style.name or ""
    if style.startswith("Heading ") and style[-1:].isdigit():
        tag = "h" + style[-1]
    alignment = {
        WD_ALIGN_PARAGRAPH.CENTER: "center",
        WD_ALIGN_PARAGRAPH.RIGHT: "right",
        WD_ALIGN_PARAGRAPH.JUSTIFY: "justify",
    }.get(paragraph.alignment, "left")
    content = f'<{tag} style="text-align:{alignment}">{"".join(parts)}</{tag}>'
    numbering = paragraph._p.pPr
    is_list = numbering is not None and numbering.numPr is not None
    if is_list or style.startswith("List"):
        ordered = "Number" in style
        if is_list:
            num_id = numbering.numPr.numId.val
            root = paragraph.part.numbering_part.element
            num = root.find(f"w:num[@w:numId='{num_id}']", root.nsmap)
            if num is not None:
                abstract_id = num.find("w:abstractNumId", root.nsmap).get(qn("w:val"))
                definition = root.find(f"w:abstractNum[@w:abstractNumId='{abstract_id}']", root.nsmap)
                if definition is not None:
                    kind = definition.find("w:lvl/w:numFmt", root.nsmap)
                    ordered = kind is not None and kind.get(qn("w:val")) != "bullet"
        tag = "ol" if ordered else "ul"
        return tag, f"<li>{''.join(parts)}</li>"
    return None, content


def _container_html(container):
    parts, list_tag = [], None
    for item in container.iter_inner_content():
        if isinstance(item, Table):
            kind, content = (
                None,
                '<table border="1">'
                + "".join(
                    "<tr>" + "".join(f"<td>{_container_html(cell)}</td>" for cell in row.cells) + "</tr>"
                    for row in item.rows
                )
                + "</table>",
            )
        else:
            kind, content = _paragraph_html(item)
        if kind != list_tag:
            if list_tag:
                parts.append(f"</{list_tag}>")
            if kind:
                parts.append(f"<{kind}>")
            list_tag = kind
        parts.append(content)
    if list_tag:
        parts.append(f"</{list_tag}>")
    return "".join(parts)


def load_docx(path):
    return _container_html(Document(path))


def _write_block(block, paragraph):
    fmt = block.blockFormat()
    level = fmt.headingLevel()
    if level:
        paragraph.style = f"Heading {min(level, 9)}"
    if block.textList():
        style = block.textList().format().style()
        bullet = style in (QTextListFormat.ListDisc, QTextListFormat.ListCircle, QTextListFormat.ListSquare)
        paragraph.style = "List Bullet" if bullet else "List Number"
    align = fmt.alignment()
    paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
        if align & Qt.AlignHCenter
        else WD_ALIGN_PARAGRAPH.RIGHT
        if align & Qt.AlignRight
        else WD_ALIGN_PARAGRAPH.JUSTIFY
        if align & Qt.AlignJustify
        else WD_ALIGN_PARAGRAPH.LEFT
    )
    paragraph.paragraph_format.left_indent = Inches(fmt.indent() * 0.25)
    iterator = block.begin()
    while not iterator.atEnd():
        fragment = iterator.fragment()
        if fragment.isValid():
            char = fragment.charFormat()
            run = paragraph.add_run()
            if char.isImageFormat():
                image = char.toImageFormat()
                source = image.name()
                if not source.startswith("data:image/"):
                    raise ValueError("Images must be embedded before DOCX export.")
                data = base64.b64decode(source.split(",", 1)[1])
                width = Inches(image.width() / 96) if image.width() > 0 else None
                run.add_picture(io.BytesIO(data), width=width)
            else:
                run.text = fragment.text()
                run.bold = char.fontWeight() >= 700
                run.italic = char.fontItalic()
                run.underline = char.fontUnderline()
                run.font.strike = char.fontStrikeOut()
                if char.fontPointSize() > 0:
                    run.font.size = Pt(char.fontPointSize())
                if char.fontFamilies():
                    run.font.name = char.fontFamilies()[0]
                if char.hasProperty(QTextCharFormat.ForegroundBrush):
                    run.font.color.rgb = RGBColor.from_string(char.foreground().color().name()[1:])
                run.font.superscript = char.verticalAlignment() == QTextCharFormat.AlignSuperScript
                run.font.subscript = char.verticalAlignment() == QTextCharFormat.AlignSubScript
                if char.isAnchor():
                    from docx.opc.constants import RELATIONSHIP_TYPE

                    relation = paragraph.part.relate_to(
                        char.anchorHref(), RELATIONSHIP_TYPE.HYPERLINK, is_external=True
                    )
                    link = OxmlElement("w:hyperlink")
                    link.set(qn("r:id"), relation)
                    link.append(run._r)
                    paragraph._p.append(link)
        iterator += 1


def _write_frame(frame, container):
    iterator = frame.begin()
    while not iterator.atEnd():
        nested = iterator.currentFrame()
        block = iterator.currentBlock()
        if isinstance(nested, QTextTable):
            table = container.add_table(rows=nested.rows(), cols=nested.columns())
            for row in range(nested.rows()):
                for column in range(nested.columns()):
                    source = nested.cellAt(row, column)
                    target = table.cell(row, column)
                    cell_block = source.firstCursorPosition().block()
                    limit = source.lastCursorPosition().position()
                    first = True
                    while cell_block.isValid() and cell_block.position() <= limit:
                        _write_block(cell_block, target.paragraphs[0] if first else target.add_paragraph())
                        first = False
                        cell_block = cell_block.next()
        elif nested is not None:
            _write_frame(nested, container)
        elif block.isValid():
            _write_block(block, container.add_paragraph())
        iterator += 1


def save_docx(path, html):
    app = QApplication.instance() or QApplication([])
    source = QTextDocument()
    source.setHtml(html)
    destination = Document()
    _write_frame(source.rootFrame(), destination)
    destination.save(path)
    del app
